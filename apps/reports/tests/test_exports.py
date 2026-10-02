"""Exportação assíncrona: XLSX e o escopo reavaliado no momento da geração."""

import io

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import AgentTeam, InternalRole, Team, TeamVisibility, User, UserKind
from apps.customers.models import Company, Plan
from apps.reports.models import ReportExport
from apps.tickets.services import ticket_service

pytestmark = pytest.mark.django_db(transaction=True)  # a task roda no on_commit


@pytest.fixture
def cenario(db):
    plan = Plan.objects.create(slug="gold", name="Gold")
    acme = Company.objects.create(name="Acme", slug="acme", plan=plan)
    suporte = Team.objects.create(name="Suporte", slug="suporte")
    financeiro = Team.objects.create(
        name="Financeiro", slug="financeiro",
        visibility=TeamVisibility.REQUESTER_ONLY,
    )
    ana = User.objects.create_user(
        email="ana@acme.test", full_name="Ana", kind=UserKind.CLIENT,
        company=acme, client_role="member",
    )
    admin = User.objects.create_user(
        email="admin@x.test", full_name="Alice", kind=UserKind.INTERNAL,
        internal_role=InternalRole.ADMIN,
    )
    agente = User.objects.create_user(
        email="agente@x.test", full_name="Samuel", kind=UserKind.INTERNAL,
        internal_role=InternalRole.AGENT,
    )
    AgentTeam.objects.create(user=agente, team=suporte)

    ticket_service.create_ticket(requester=ana, team=suporte,
                                 subject="No suporte", body="b")
    ticket_service.create_ticket(requester=ana, team=financeiro,
                                 subject="No financeiro", body="b")
    return {
        "admin": admin, "agente": agente, "ana": ana,
        "suporte": suporte, "financeiro": financeiro,
    }


@pytest.fixture
def api():
    def _for(user):
        client = APIClient()
        client.force_authenticate(user=user)
        return client
    return _for


def test_xlsx_export_completes(api, cenario):
    resp = api(cenario["admin"]).post(
        "/api/v1/reports/exports/", {"report_key": "tickets", "format": "xlsx"},
        format="json",
    )
    assert resp.status_code == 201

    export = ReportExport.objects.get(pk=resp.json()["id"])
    assert export.status == ReportExport.Status.DONE
    assert export.row_count == 2
    assert export.expires_at is not None       # D18: expurgo programado
    assert export.file.name.endswith(".xlsx")

    download = api(cenario["admin"]).get(f"/api/v1/reports/exports/{export.pk}/download/")
    assert download.status_code == 200
    content = b"".join(download.streaming_content)
    assert content[:2] == b"PK"                # xlsx é um zip


def test_xlsx_export_respects_team_filter(api, cenario):
    """Espelha `test_csv_respects_team_filter` (síncrono) para o canal
    assíncrono: pedir XLSX com o painel filtrado por equipe não pode voltar
    a baixar as demais equipes que o usuário enxerga — mesmo bug do CSV,
    canal diferente (`build_report_export` monta seus próprios filtros a
    partir de `export.params` e não lia `team`)."""
    resp = api(cenario["admin"]).post(
        "/api/v1/reports/exports/",
        {
            "report_key": "tickets", "format": "xlsx",
            "params": {"team": cenario["suporte"].pk},
        },
        format="json",
    )
    assert resp.status_code == 201

    export = ReportExport.objects.get(pk=resp.json()["id"])
    assert export.status == ReportExport.Status.DONE, export.error
    assert export.row_count == 1  # só o ticket de Suporte, não o do Financeiro

    download = api(cenario["admin"]).get(f"/api/v1/reports/exports/{export.pk}/download/")
    wb = load_workbook(io.BytesIO(b"".join(download.streaming_content)))
    valores = [cell.value for row in wb["Tickets"].iter_rows(min_row=2) for cell in row]
    assert "Suporte" in valores
    assert "Financeiro" not in valores


def test_scope_is_the_requesters(api, cenario):
    """Agente de Suporte exporta só o que ele enxerga — não o quadro todo."""
    resp = api(cenario["agente"]).post(
        "/api/v1/reports/exports/", {"format": "csv"}, format="json"
    )
    export = ReportExport.objects.get(pk=resp.json()["id"])
    assert export.row_count == 1               # sem o ticket do Financeiro


def test_client_cannot_export(api, cenario):
    assert api(cenario["ana"]).post(
        "/api/v1/reports/exports/", {"format": "csv"}, format="json"
    ).status_code == 403


def test_user_sees_only_own_exports(api, cenario):
    api(cenario["admin"]).post("/api/v1/reports/exports/", {"format": "csv"},
                               format="json")
    assert api(cenario["agente"]).get("/api/v1/reports/exports/").json()["count"] == 0


def test_download_before_ready_is_404(api, cenario):
    export = ReportExport.objects.create(
        requested_by=cenario["admin"], report_key="tickets", format="csv",
    )
    assert api(cenario["admin"]).get(
        f"/api/v1/reports/exports/{export.pk}/download/"
    ).status_code == 404


def test_pdf_export_completes(api, cenario):
    """O caminho do PDF ficou meses sem teste porque o WeasyPrint era extra
    opcional não-instalado — a suíte inteira passava com a exportação quebrada.
    Agora o extra faz parte do ambiente de teste e o caminho fica coberto."""
    resp = api(cenario["admin"]).post(
        "/api/v1/reports/exports/", {"report_key": "tickets", "format": "pdf"},
        format="json",
    )
    assert resp.status_code == 201

    export = ReportExport.objects.get(pk=resp.json()["id"])
    assert export.status == ReportExport.Status.DONE, export.error
    assert export.file.name.endswith(".pdf")

    download = api(cenario["admin"]).get(
        f"/api/v1/reports/exports/{export.pk}/download/"
    )
    assert download.status_code == 200
    content = b"".join(download.streaming_content)
    assert content[:5] == b"%PDF-"  # é um PDF de verdade, não um HTML com outro nome
